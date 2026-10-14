"""Quick forecast infers the season from the data (audit finding C-05).

Before: the workbench always sent season 7 + ETS, and the API had no way to
infer a cycle, so persona C's Q4-peaked monthly series forecast flat at about
2,075 (January 2,076 against last January's 1,040). Now ``season_length`` and
``method_id`` accept ``"auto"``; the response says which settings were used.
"""

from __future__ import annotations

import pytest
from fastapi.testclient import TestClient

from demand_signal_os.api.app import create_app

MARCO_MONTHLY = [
    980, 940, 1050, 1100, 1180, 1120, 1060, 1090, 1250, 1420, 1780, 1950,
    1040, 990, 1110, 1170, 1240, 1190, 1130, 1160, 1330, 1510, 1890, 2080,
]

_AUTO = {
    "history": MARCO_MONTHLY,
    "horizon": 12,
    "season_length": "auto",
    "method_id": "auto",
    "band": True,
}


@pytest.fixture(scope="module")
def client() -> TestClient:
    return TestClient(create_app())


def test_auto_detects_yearly_cycle_and_reports_it(client: TestClient) -> None:
    resp = client.post("/api/v1/forecast/single", json=_AUTO)
    assert resp.status_code == 200, resp.text
    settings = resp.json()["settings"]
    assert settings["season_length"] == 12
    assert settings["season_source"] == "detected"
    assert settings["method_source"] == "auto"
    assert "12-period cycle" in settings["note"]


def test_auto_forecast_follows_the_q4_peak_instead_of_flat(client: TestClient) -> None:
    band = client.post("/api/v1/forecast/single", json=_AUTO).json()["band"]
    q50 = [step["q50"] for step in band]
    # Month 1 of the forecast is a January (history starts in a January):
    # it must sit near the January actuals, not at the December peak level.
    assert q50[0] < 1300, q50
    # The forecast December (step 12) is the peak, well above January.
    assert q50[11] > 1.6 * q50[0], q50
    assert max(q50) == q50[11]


def test_explicit_season_still_gives_the_old_answer(client: TestClient) -> None:
    # Backward compatible: an explicit 7 + ets is untouched (no detection).
    req = {**_AUTO, "season_length": 7, "method_id": "ets"}
    body = client.post("/api/v1/forecast/single", json=req).json()
    assert body["settings"]["season_length"] == 7
    assert body["settings"]["season_source"] == "given"
    assert body["method"] == "ets"


def test_exports_use_the_same_resolved_settings(client: TestClient) -> None:
    import io

    from openpyxl import load_workbook

    resp = client.post("/api/v1/forecast/single.xlsx", json=_AUTO)
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content))
    rows = {r[0]: r[1] for r in wb["Provenance"].iter_rows(values_only=True) if r and r[0]}
    assert rows.get("settings.season_length") == 12
    assert rows.get("settings.season_source") == "detected"


# Fixed-seed baseline (RULE 5). Pinned from the first observed run on
# 2026-10-13; drift means a non-deterministic component landed in the
# detect -> resolve -> fit -> sample path.
_BASELINE_METHOD = "ets"  # ETS fits the 12-cycle on this series (ETS(M,N,M))
_BASELINE_Q50 = [
    1102.037, 1052.754, 1177.689, 1236.875, 1318.226, 1257.713,
    1191.845, 1223.875, 1403.019, 1592.583, 1994.148, 2188.907,
]


def test_auto_forecast_fixed_seed_baseline(client: TestClient) -> None:
    a = client.post("/api/v1/forecast/single", json=_AUTO).json()
    b = client.post("/api/v1/forecast/single", json=_AUTO).json()
    assert a["band"] == b["band"]
    assert a["settings"] == b["settings"]
    assert a["method"] == _BASELINE_METHOD
    assert [round(s["q50"], 3) for s in a["band"]] == _BASELINE_Q50
