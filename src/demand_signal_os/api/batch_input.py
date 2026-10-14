"""Parse a multi-SKU demand file into per-(SKU, location) series.

Audit findings C-08 / F-03: a planner with 45 SKU x DC series could only paste
one series at a time, and on plan2cash.com "Run forecast" ignored a validated
demand_history workbook. This module turns ONE uploaded file into the list of
series the batch forecast route fits.

Accepted shapes (detected, never guessed silently):

* the ``demand_history_multi`` template (``observations`` sheet: sku_id,
  location_id, period_label, observed_demand, censoring_flag);
* the single-SKU ``demand_history`` template (``Slots`` F5/F6 + ``History``);
* any CSV or xlsx table with a header row naming a SKU column, a demand column
  and optionally a location and a period column. Header names are matched
  case-insensitively against the aliases below.

Rows keep their file order unless every period cell is a date, in which case
each series is sorted by date. Same file in, same series out (RULE 5).
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 50_000
MAX_SERIES = 200

_SKU_ALIASES = {"sku_id", "sku", "sku_identifier", "item", "item_id", "product", "product_id", "material"}
_LOC_ALIASES = {"location_id", "location", "location_identifier", "dc", "dc_id", "site", "warehouse", "store"}
_PERIOD_ALIASES = {"period_label", "period", "date", "week", "month", "day", "bucket"}
_DEMAND_ALIASES = {
    "observed_demand", "demand", "units", "qty", "quantity", "sales", "units_sold", "volume", "orders",
}

DEFAULT_LOCATION = "ALL"


class BatchInputError(ValueError):
    """A file we cannot read. The message is plain and safe to show a user."""


@dataclass
class ParsedSeries:
    sku_id: str
    location_id: str
    periods: list[str] = field(default_factory=list)
    values: list[float] = field(default_factory=list)


@dataclass
class ParsedBatch:
    source_format: str
    series: list[ParsedSeries]
    skipped_rows: int = 0


def _norm(header: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(header or "").strip().lower()).strip("_")


def _find(headers: list[str], aliases: set[str]) -> int | None:
    for i, h in enumerate(headers):
        if h in aliases:
            return i
    return None


def _to_float(raw: Any) -> float | None:
    if raw is None:
        return None
    if isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        value = float(raw)
    else:
        text = str(raw).strip().replace(",", "")
        if not text:
            return None
        try:
            value = float(text)
        except ValueError:
            return None
    return value if value == value and value not in (float("inf"), float("-inf")) else None


def _period_text(raw: Any) -> str:
    if isinstance(raw, datetime):
        return raw.date().isoformat()
    if isinstance(raw, date):
        return raw.isoformat()
    return "" if raw is None else str(raw).strip()


def _as_date(text: str) -> date | None:
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _group(rows: list[tuple[str, str, str, float]], source_format: str, skipped: int) -> ParsedBatch:
    """Group (sku, location, period, value) rows into series, file order kept."""
    grouped: dict[tuple[str, str], ParsedSeries] = {}
    for sku, loc, period, value in rows:
        key = (sku, loc)
        if key not in grouped:
            if len(grouped) >= MAX_SERIES:
                raise BatchInputError(
                    f"The file has more than {MAX_SERIES} SKU and location series. "
                    f"Split it into files of up to {MAX_SERIES} series."
                )
            grouped[key] = ParsedSeries(sku_id=sku, location_id=loc)
        grouped[key].periods.append(period)
        grouped[key].values.append(value)
    for s in grouped.values():
        dates = [_as_date(p) for p in s.periods]
        if s.periods and all(d is not None for d in dates):
            order = sorted(range(len(dates)), key=lambda i: (dates[i], i))
            s.periods = [s.periods[i] for i in order]
            s.values = [s.values[i] for i in order]
    if not grouped:
        raise BatchInputError("No demand rows were found in the file.")
    return ParsedBatch(source_format=source_format, series=list(grouped.values()), skipped_rows=skipped)


def _from_table(table: list[list[Any]], source_format: str, default_sku: str | None = None,
                default_loc: str | None = None) -> ParsedBatch:
    """Header row + data rows -> grouped series (generic CSV/xlsx table)."""
    # The header is the first row that names a demand column (templates and
    # exports sometimes carry a title row above it).
    header_idx = None
    for i, row in enumerate(table[:10]):
        if _find([_norm(c) for c in row], _DEMAND_ALIASES) is not None:
            header_idx = i
            break
    if header_idx is None:
        raise BatchInputError(
            "Could not find a demand column. Use a header row with columns such as "
            "sku_id, location_id, period and units (or demand)."
        )
    headers = [_norm(c) for c in table[header_idx]]
    i_val = _find(headers, _DEMAND_ALIASES)
    i_sku = _find(headers, _SKU_ALIASES)
    i_loc = _find(headers, _LOC_ALIASES)
    i_per = _find(headers, _PERIOD_ALIASES)
    if i_sku is None and default_sku is None:
        raise BatchInputError(
            "Could not find a SKU column. Add a column named sku_id (or sku) so each "
            "row says which item it belongs to."
        )
    rows: list[tuple[str, str, str, float]] = []
    skipped = 0
    for n, row in enumerate(table[header_idx + 1:], start=1):
        if n > MAX_ROWS:
            raise BatchInputError(f"The file has more than {MAX_ROWS:,} rows. Split it into smaller files.")
        if not any(c not in (None, "") for c in row):
            continue

        def cell(i: int | None, _row: list[Any] = row) -> Any:
            return _row[i] if i is not None and i < len(_row) else None

        value = _to_float(cell(i_val))
        sku = str(cell(i_sku) or default_sku or "").strip()
        if value is None or not sku:
            skipped += 1
            continue
        loc = str(cell(i_loc) or default_loc or DEFAULT_LOCATION).strip() or DEFAULT_LOCATION
        rows.append((sku, loc, _period_text(cell(i_per)), value))
    return _group(rows, source_format, skipped)


def _read_csv(data: bytes) -> list[list[Any]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    sample = text[:4096]
    delimiter = ","
    if sample.count(";") > sample.count(",") and sample.count(";") >= sample.count("\t"):
        delimiter = ";"
    elif sample.count("\t") > sample.count(","):
        delimiter = "\t"
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _sheet_rows(ws: Any) -> list[list[Any]]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def parse_batch_file(data: bytes) -> ParsedBatch:
    """Parse an uploaded CSV or xlsx into per-(SKU, location) series."""
    if not data:
        raise BatchInputError("The file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise BatchInputError("The file is larger than 5 MB. Split it into smaller files.")
    if not data.startswith(b"PK"):
        return _from_table(_read_csv(data), "csv")

    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any unreadable workbook
        raise BatchInputError("The file could not be opened as an Excel workbook (.xlsx).") from exc
    try:
        names = {n.lower(): n for n in wb.sheetnames}
        if "observations" in names:
            return _from_table(_sheet_rows(wb[names["observations"]]), "demand_history_multi")
        if "history" in names and "slots" in names:
            slots = wb[names["slots"]]
            sku = str(slots["F5"].value or "").strip() or "SKU-1"
            loc = str(slots["F6"].value or "").strip() or DEFAULT_LOCATION
            return _from_table(_sheet_rows(wb[names["history"]]), "demand_history",
                               default_sku=sku, default_loc=loc)
        return _from_table(_sheet_rows(wb[wb.sheetnames[0]]), "xlsx")
    finally:
        wb.close()


__all__ = [
    "MAX_FILE_BYTES",
    "MAX_ROWS",
    "MAX_SERIES",
    "BatchInputError",
    "ParsedBatch",
    "ParsedSeries",
    "parse_batch_file",
]
